import asyncio
import csv
import json
import logging
import os
import subprocess
import sys
import tempfile
import threading
import time
import uuid
from datetime import datetime
from pathlib import Path
from typing import Optional, List

import httpx

# Force UTF-8 everywhere to handle Unicode transcription output
os.environ.setdefault("PYTHONIOENCODING", "utf-8")
# Disable torchcodec - we don't need it and it requires FFmpeg shared DLLs on Windows
os.environ["TORCHCODEC_DISABLE"] = "1"
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

import torchaudio
import torch
from docx import Document
from docx.shared import Pt
from fastapi import FastAPI, File, HTTPException, UploadFile, Request, Form
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from stt.gigaam_model import gigaam_model, set_num_threads
from stt.whisper_model import whisper_small_model, whisper_medium_model
from stt.gigaam_multilingual_model import gigaam_multilingual_model

AVAILABLE_MODELS = {
    "GigaAM-v3": gigaam_model,
    "GigaAM Multilingual": gigaam_multilingual_model,
    "Whisper Small": whisper_small_model,
    "Whisper Medium": whisper_medium_model,
}

def get_model(name: str = "GigaAM-v3"):
    return AVAILABLE_MODELS.get(name, gigaam_model)

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    encoding='utf-8' if sys.version_info >= (3, 9) else None
)
logger = logging.getLogger(__name__)

app = FastAPI(
    title="Orpheus",
    version="1.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

MAX_FILE_SIZE_MB = 100
MAX_CHUNK_DURATION_SEC = 25
CHUNK_OVERLAP_SEC = 1
OUTPUT_DIR = Path("./output")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

LOG_FILE = Path("orphrus.log")
LOG_HEADER = ["number", "date", "file_name", "duration", "file_size", "run_time", "num_chunks", "cpu_cores", "model"]

stop_events = {}
transcription_locks = {}
transcription_progress = {}
batch_progress = {}
batch_results = {}

if not LOG_FILE.exists():
    with open(LOG_FILE, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f, delimiter='\t')
        writer.writerow(LOG_HEADER)
else:
    # Convert existing log file to UTF-8 if not already
    for enc in ('utf-8', 'cp1251', 'latin-1'):
        try:
            with open(LOG_FILE, 'r', encoding=enc) as f:
                content = f.read()
            with open(LOG_FILE, 'w', newline='', encoding='utf-8') as f:
                f.write(content)
            break
        except UnicodeDecodeError:
            continue


class TranscriptionResponse(BaseModel):
    text: str
    success: bool
    error: Optional[str] = None
    audio_info: Optional[dict] = None
    runtime_seconds: Optional[float] = None
    runtime_formatted: Optional[str] = None
    progress: Optional[int] = None
    transcription_id: Optional[str] = None


class ProgressResponse(BaseModel):
    current_chunk: int
    total_chunks: int
    status: str


class SaveRequest(BaseModel):
    text: str
    output_path: str


class SaveResponse(BaseModel):
    success: bool
    message: str
    file_path: Optional[str] = None


class ModelInfoResponse(BaseModel):
    loaded: bool
    model_type: str
    revision: str


class LLMEnhanceRequest(BaseModel):
    text: str
    prompt: str
    model: Optional[str] = "minimax-m3:cloud"
    ollama_host: str = "http://localhost:11434"


class ChunkRequest(BaseModel):
    start_sec: float
    end_sec: Optional[float] = None


class TranscribeChunksRequest(BaseModel):
    filename: str
    chunks: List[ChunkRequest]


class ChunkTranscriptionResult(BaseModel):
    chunk_index: int
    text: str
    success: bool
    error: Optional[str] = None


class TranscribeChunksResponse(BaseModel):
    results: List[ChunkTranscriptionResult]
    success: bool
    total_chunks: int
    processed_chunks: int
    runtime_seconds: Optional[float] = None
    runtime_formatted: Optional[str] = None
    error: Optional[str] = None


SUPPORTED_FORMATS = ('.mp3', '.wav', '.flac', '.m4a', '.ogg', '.aiff', '.aif', '.caf', '.mp4')


def get_log_count():
    if not LOG_FILE.exists():
        return 0
    for enc in ('utf-8', 'cp1251', 'latin-1'):
        try:
            with open(LOG_FILE, 'r', encoding=enc) as f:
                lines = [l for l in f.readlines() if l.strip() and not l.startswith('number')]
                return len(lines)
        except UnicodeDecodeError:
            continue
    return 0


def save_log_entry(file_name, duration, file_size_mb, run_time, num_chunks, cpu_cores, model_name="GigaAM-v3"):
    number = get_log_count() + 1
    date = datetime.now().strftime('%Y-%m-%d')
    hours = int(duration // 3600)
    mins = int((duration % 3600) // 60)
    secs = int(duration % 60)
    duration_str = f"{hours}:{mins:02d}:{secs:02d}" if hours > 0 else f"{mins:02d}:{secs:02d}"
    hours_r = int(run_time // 3600)
    mins_r = int((run_time % 3600) // 60)
    secs_r = int(run_time % 60)
    runtime_formatted = f"{hours_r:02d}:{mins_r:02d}:{secs_r:02d}"
    file_size_str = f"{file_size_mb:.1f} Mb"
    with open(LOG_FILE, 'a', newline='', encoding='utf-8') as f:
        writer = csv.writer(f, delimiter='\t')
        writer.writerow([number, date, file_name, duration_str, file_size_str, runtime_formatted, num_chunks, cpu_cores, model_name])
    logger.info(f"Log entry #{number} saved")


def validate_file_format(filename: Optional[str]) -> bool:
    if not filename:
        return False
    return filename.lower().endswith(SUPPORTED_FORMATS)


def get_audio_info(file_path: str, file_size: int) -> dict:
    info = torchaudio.info(file_path)
    duration_sec = info.num_frames / info.sample_rate
    hours = int(duration_sec // 3600)
    minutes = int((duration_sec % 3600) // 60)
    seconds = int(duration_sec % 60)
    return {
        "duration_formatted": f"{hours:02d}:{minutes:02d}:{seconds:02d}",
        "duration_seconds": round(duration_sec, 2),
        "channels": info.num_channels,
        "sample_rate": info.sample_rate,
        "file_size_mb": round(file_size / (1024 * 1024), 2)
    }


def calculate_num_chunks(duration_seconds: float) -> int:
    if duration_seconds <= MAX_CHUNK_DURATION_SEC:
        return 1
    effective_chunk = MAX_CHUNK_DURATION_SEC - CHUNK_OVERLAP_SEC
    num_chunks = int((duration_seconds - MAX_CHUNK_DURATION_SEC) / effective_chunk) + 2
    return max(1, num_chunks)


def split_audio_into_chunks(audio_path: str, num_chunks: int) -> List[str]:
    waveform, sample_rate = torchaudio.load(audio_path)
    if sample_rate != 16000:
        waveform = torchaudio.functional.resample(waveform, sample_rate, 16000)
        sample_rate = 16000
    total_samples = waveform.shape[1]
    chunk_samples = MAX_CHUNK_DURATION_SEC * sample_rate
    hop_samples = (MAX_CHUNK_DURATION_SEC - CHUNK_OVERLAP_SEC) * sample_rate
    temp_dir = os.path.dirname(audio_path) or "."
    temp_prefix = os.path.splitext(os.path.basename(audio_path))[0]
    chunks = []
    for i in range(num_chunks):
        start = i * hop_samples
        end = min(start + chunk_samples, total_samples)
        chunk_waveform = waveform[:, start:end]
        chunk_path = os.path.join(temp_dir, f"{temp_prefix}_chunk_{i}.wav")
        torchaudio.save(chunk_path, chunk_waveform, sample_rate)
        chunks.append(chunk_path)
        if end >= total_samples:
            break
    return chunks


def cleanup_files(files: List[Optional[str]]) -> None:
    for file_path in files:
        try:
            if file_path and os.path.exists(file_path):
                os.unlink(file_path)
        except Exception as e:
            logger.warning(f"Failed to cleanup {file_path}: {e}")


def convert_to_wav(input_path: str) -> str:
    if input_path.lower().endswith('.wav'):
        return input_path
    output_path = os.path.splitext(input_path)[0] + '.wav'
    try:
        result = subprocess.run(
            ['ffmpeg', '-y', '-i', input_path, '-ar', '16000', '-ac', '1', output_path],
            capture_output=True, text=True
        )
        if result.returncode != 0:
            raise RuntimeError(f"FFmpeg conversion failed: {result.stderr}")
    except FileNotFoundError:
        raise RuntimeError("FFmpeg not found. Please install FFmpeg and add it to PATH.")
    except Exception as e:
        raise RuntimeError(f"Failed to convert audio to WAV: {e}")
    try:
        os.unlink(input_path)
    except Exception:
        pass
    return output_path


def extract_chunk(audio_path: str, start_sec: float, end_sec: Optional[float] = None) -> str:
    waveform, sample_rate = torchaudio.load(audio_path)
    if sample_rate != 16000:
        waveform = torchaudio.functional.resample(waveform, sample_rate, 16000)
        sample_rate = 16000
    
    total_samples = waveform.shape[1]
    total_duration = total_samples / sample_rate
    
    if end_sec is None:
        end_sec = total_duration
    
    start_sample = int(start_sec * sample_rate)
    end_sample = int(end_sec * sample_rate)
    end_sample = min(end_sample, total_samples)
    
    chunk_waveform = waveform[:, start_sample:end_sample]
    
    temp_dir = os.path.dirname(audio_path) or "."
    chunk_path = os.path.join(temp_dir, f"chunk_{start_sec}_{end_sec}.wav")
    torchaudio.save(chunk_path, chunk_waveform, sample_rate)
    return chunk_path


@app.on_event("startup")
async def startup_event():
    logger.info("Starting Orpheus API...")
    os.environ["HF_HUB_DOWNLOAD_TIMEOUT"] = "120"
    set_num_threads(1)


@app.get("/")
async def root():
    return {"message": "Orpheus API", "version": "1.0.0"}


@app.get("/api/model-info", response_model=ModelInfoResponse)
async def get_model_info():
    info = gigaam_model.get_model_info()
    return ModelInfoResponse(**info)


@app.get("/api/models")
async def list_models():
    return {
        "models": [
            {"id": "GigaAM-v3", "name": "GigaAM-v3", "loaded": gigaam_model._model is not None},
            {"id": "GigaAM Multilingual", "name": "GigaAM Multilingual", "loaded": gigaam_multilingual_model._model is not None},
            {"id": "Whisper Small", "name": "Whisper Small", "loaded": whisper_small_model._pipeline is not None},
            {"id": "Whisper Medium", "name": "Whisper Medium", "loaded": whisper_medium_model._pipeline is not None},
        ]
    }


@app.post("/api/settings")
async def update_settings(num_threads: int = 1):
    num_threads = max(1, min(4, num_threads))
    set_num_threads(num_threads)
    logger.info(f"CPU threads set to: {num_threads}")
    return {"success": True, "cpu_threads": num_threads}


@app.get("/api/settings")
async def get_settings():
    return {"cpu_threads": torch.get_num_threads()}


@app.post("/api/stop")
async def stop_transcription(tid: str = ""):
    logger.info(f"Stop requested for: {tid}")
    if tid:
        if tid in stop_events:
            stop_events[tid].set()
    else:
        for se in stop_events.values():
            se.set()
    return {"success": True, "message": "Stop requested"}


@app.get("/api/progress/{tid}")
async def get_progress(tid: str):
    if tid in transcription_progress:
        return ProgressResponse(
            current_chunk=transcription_progress[tid].get("current_chunk", 0),
            total_chunks=transcription_progress[tid].get("total_chunks", 0),
            status=transcription_progress[tid].get("status", "processing")
        )
    return ProgressResponse(current_chunk=0, total_chunks=0, status="not_found")


@app.post("/api/audio-info")
async def get_audio_info_endpoint(file: UploadFile = File(...)):
    if not validate_file_format(file.filename):
        return {"success": False, "error": "Unsupported file format"}
    
    temp_audio_path = None
    try:
        content = await file.read()
        file_size_mb = len(content) / (1024 * 1024)
        if file_size_mb > MAX_FILE_SIZE_MB:
            return {"success": False, "error": f"File too large. Max {MAX_FILE_SIZE_MB}MB"}
        
        suffix = Path(file.filename).suffix
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
            temp_file.write(content)
            temp_audio_path = temp_file.name
        
        temp_audio_path = convert_to_wav(temp_audio_path)
        
        audio_info = get_audio_info(temp_audio_path, len(content))
        duration = audio_info["duration_seconds"]
        num_chunks = calculate_num_chunks(duration)
        audio_info["num_chunks"] = num_chunks
        audio_info["max_chunk_duration"] = MAX_CHUNK_DURATION_SEC
        return {"success": True, "audio_info": audio_info}
    except Exception as e:
        return {"success": False, "error": str(e)}
    finally:
        cleanup_files([temp_audio_path])


def run_transcription(transcription_id: str, content: bytes, filename: str, result_container: dict, model_name: str = "GigaAM-v3"):
    """Run transcription in a separate thread"""
    model = get_model(model_name)
    stop_event = stop_events.get(transcription_id)
    temp_audio_path = None
    chunk_paths = []
    start_time = time.time()
    
    transcription_progress[transcription_id] = {"current_chunk": 0, "total_chunks": 0, "status": "starting"}
    
    try:
        file_size_mb = len(content) / (1024 * 1024)
        
        suffix = Path(filename).suffix
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
            temp_file.write(content)
            temp_audio_path = temp_file.name
        
        temp_audio_path = convert_to_wav(temp_audio_path)
        
        if stop_event and stop_event.is_set():
            result_container["result"] = {"text": "", "success": False, "error": "Transcription stopped by user"}
            transcription_progress[transcription_id] = {"current_chunk": 0, "total_chunks": 0, "status": "stopped"}
            return
        
        audio_info = get_audio_info(temp_audio_path, len(content))
        duration = audio_info["duration_seconds"]
        num_chunks = calculate_num_chunks(duration)
        
        transcription_progress[transcription_id] = {"current_chunk": 0, "total_chunks": num_chunks, "status": "processing"}
        
        if num_chunks > 1:
            chunk_paths = split_audio_into_chunks(temp_audio_path, num_chunks)
            logger.info(f"Split into {len(chunk_paths)} chunks")
            texts = []
            for i, chunk_path in enumerate(chunk_paths):
                if stop_event and stop_event.is_set():
                    result_container["result"] = {"text": "", "success": False, "error": "Transcription stopped by user"}
                    transcription_progress[transcription_id] = {"current_chunk": i, "total_chunks": len(chunk_paths), "status": "stopped"}
                    return
                logger.info(f"Processing chunk {i+1}/{len(chunk_paths)}...")
                transcription_progress[transcription_id] = {"current_chunk": i + 1, "total_chunks": len(chunk_paths), "status": "processing"}
                chunk_text = model.transcribe(chunk_path)
                texts.append(chunk_text)
            text = " ".join(texts)
        else:
            if stop_event and stop_event.is_set():
                result_container["result"] = {"text": "", "success": False, "error": "Transcription stopped by user"}
                transcription_progress[transcription_id] = {"current_chunk": 0, "total_chunks": 1, "status": "stopped"}
                return
            transcription_progress[transcription_id] = {"current_chunk": 1, "total_chunks": 1, "status": "processing"}
            text = model.transcribe(temp_audio_path)
        
        if stop_event and stop_event.is_set():
            result_container["result"] = {"text": "", "success": False, "error": "Transcription stopped by user"}
            transcription_progress[transcription_id] = {"current_chunk": 0, "total_chunks": num_chunks, "status": "stopped"}
            return
        
        runtime_seconds = time.time() - start_time
        minutes = int(runtime_seconds // 60)
        seconds = int(runtime_seconds % 60)
        runtime_formatted = f"{minutes:02d}:{seconds:02d}"
        cpu_cores = torch.get_num_threads()
        save_log_entry(filename, duration, file_size_mb, runtime_seconds, num_chunks, cpu_cores, model_name)
        
        result_container["result"] = {
            "text": text,
            "success": True,
            "audio_info": audio_info,
            "runtime_seconds": round(runtime_seconds, 2),
            "runtime_formatted": runtime_formatted
        }
        transcription_progress[transcription_id] = {"current_chunk": num_chunks, "total_chunks": num_chunks, "status": "completed", "result": result_container["result"]}
    except Exception as e:
        logger.exception("Transcription failed")
        result_container["result"] = {"text": "", "success": False, "error": str(e)}
        transcription_progress[transcription_id] = {"current_chunk": 0, "total_chunks": 0, "status": "error", "result": {"error": str(e)}}
    finally:
        cleanup_files(chunk_paths + [temp_audio_path])
        if transcription_id in stop_events:
            del stop_events[transcription_id]


class TranscriptionStartResponse(BaseModel):
    transcription_id: str
    success: bool
    error: Optional[str] = None


class BatchStartResponse(BaseModel):
    batch_id: str
    success: bool
    error: Optional[str] = None


class BatchProgressResponse(BaseModel):
    files_total: int
    files_done: int
    current_file: str
    current_chunk: int
    total_chunks: int
    status: str


class BatchFileResult(BaseModel):
    filename: str
    text: str
    success: bool
    error: Optional[str] = None
    runtime_seconds: Optional[float] = None
    doc_path: Optional[str] = None


class BatchResultResponse(BaseModel):
    success: bool
    combined_text: str
    files: List[BatchFileResult]
    total_runtime_seconds: Optional[float] = None
    total_runtime_formatted: Optional[str] = None
    error: Optional[str] = None


@app.post("/api/transcribe-start", response_model=TranscriptionStartResponse)
async def start_transcription(file: UploadFile = File(...), model: str = Form("GigaAM-v3")):
    if not validate_file_format(file.filename):
        return TranscriptionStartResponse(transcription_id="", success=False, error="Unsupported file format")
    
    content = await file.read()
    file_size_mb = len(content) / (1024 * 1024)
    if file_size_mb > MAX_FILE_SIZE_MB:
        return TranscriptionStartResponse(transcription_id="", success=False, error=f"File too large. Max {MAX_FILE_SIZE_MB}MB")
    
    transcription_id = str(uuid.uuid4())
    stop_event = threading.Event()
    stop_events[transcription_id] = stop_event
    
    result_container = {}
    
    thread = threading.Thread(target=run_transcription, args=(transcription_id, content, file.filename, result_container, model))
    thread.start()
    
    return TranscriptionStartResponse(transcription_id=transcription_id, success=True)


@app.get("/api/transcribe-result/{tid}")
async def get_transcription_result(tid: str):
    if tid not in transcription_progress:
        return TranscriptionResponse(text="", success=False, error="Transcription not found")
    
    progress_info = transcription_progress[tid]
    status = progress_info.get("status", "")
    
    if status == "completed":
        result = progress_info.get("result", {})
        transcription_progress.pop(tid, None)
        return TranscriptionResponse(
            text=result.get("text", ""),
            success=True,
            audio_info=result.get("audio_info"),
            runtime_seconds=result.get("runtime_seconds"),
            runtime_formatted=result.get("runtime_formatted")
        )
    elif status == "error" or status == "stopped":
        result = progress_info.get("result", {})
        transcription_progress.pop(tid, None)
        return TranscriptionResponse(text="", success=False, error=result.get("error", "Unknown error"))
    
    return TranscriptionResponse(text="", success=False, error="Transcription in progress")


@app.post("/api/transcribe-chunks", response_model=TranscribeChunksResponse)
async def transcribe_chunks(file: UploadFile = File(...), chunks_json: str = Form(""), model: str = Form("GigaAM-v3")):
    if not validate_file_format(file.filename):
        return TranscribeChunksResponse(
            results=[], success=False, total_chunks=0, processed_chunks=0,
            error="Unsupported file format"
        )
    
    content = await file.read()
    file_size_mb = len(content) / (1024 * 1024)
    if file_size_mb > MAX_FILE_SIZE_MB:
        return TranscribeChunksResponse(
            results=[], success=False, total_chunks=0, processed_chunks=0,
            error=f"File too large. Max {MAX_FILE_SIZE_MB}MB"
        )
    
    try:
        chunks_data = json.loads(chunks_json) if chunks_json else []
    except Exception as e:
        logger.error(f"JSON parse error: {e}")
        chunks_data = []
    
    mdl = get_model(model)
    temp_audio_path = None
    chunk_paths = []
    results = []
    processed = 0
    start_time = time.time()
    
    try:
        suffix = Path(file.filename).suffix
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
            temp_file.write(content)
            temp_audio_path = temp_file.name
        
        temp_audio_path = convert_to_wav(temp_audio_path)
        audio_info = get_audio_info(temp_audio_path, len(content))
        total_duration = audio_info["duration_seconds"]
        
        for i, chunk in enumerate(chunks_data):
            try:
                start_sec = chunk.get("start_sec", 0)
                end_sec = chunk.get("end_sec")
                if end_sec is None:
                    end_sec = total_duration
                
                chunk_duration = end_sec - start_sec
                
                if chunk_duration <= MAX_CHUNK_DURATION_SEC:
                    chunk_path = extract_chunk(temp_audio_path, start_sec, end_sec)
                    chunk_paths.append(chunk_path)
                    text = mdl.transcribe(chunk_path)
                    results.append(ChunkTranscriptionResult(
                        chunk_index=i,
                        text=text,
                        success=True
                    ))
                    processed += 1
                else:
                    sub_chunks = calculate_num_chunks(chunk_duration)
                    sub_chunk_duration = chunk_duration / sub_chunks
                    sub_texts = []
                    for sc in range(sub_chunks):
                        sub_start = start_sec + sc * sub_chunk_duration
                        sub_end = min(sub_start + sub_chunk_duration, end_sec)
                        sub_chunk_path = extract_chunk(temp_audio_path, sub_start, sub_end)
                        chunk_paths.append(sub_chunk_path)
                        sub_text = mdl.transcribe(sub_chunk_path)
                        sub_texts.append(sub_text)
                    results.append(ChunkTranscriptionResult(
                        chunk_index=i,
                        text=" ".join(sub_texts),
                        success=True
                    ))
                    processed += 1
            except Exception as e:
                results.append(ChunkTranscriptionResult(
                    chunk_index=i,
                    text="",
                    success=False,
                    error=str(e)
                ))
        
        runtime_seconds = time.time() - start_time
        minutes = int(runtime_seconds // 60)
        seconds = int(runtime_seconds % 60)
        runtime_formatted = f"{minutes:02d}:{seconds:02d}"
        
        total_duration = sum(ch.get("end_sec", 0) - ch.get("start_sec", 0) for ch in chunks_data)
        save_log_entry(file.filename, total_duration, file_size_mb, runtime_seconds, len(chunks_data), torch.get_num_threads(), model)
        
        return TranscribeChunksResponse(
            results=results,
            success=True,
            total_chunks=len(chunks_data),
            processed_chunks=processed,
            runtime_seconds=round(runtime_seconds, 2),
            runtime_formatted=runtime_formatted
        )
    except Exception as e:
        return TranscribeChunksResponse(
            results=results,
            success=False,
            total_chunks=len(chunks_data),
            processed_chunks=processed,
            error=str(e)
        )
    finally:
        cleanup_files(chunk_paths + [temp_audio_path])


@app.post("/api/transcribe", response_model=TranscriptionResponse)
async def transcribe_audio(file: UploadFile = File(...)):
    if not validate_file_format(file.filename):
        return TranscriptionResponse(text="", success=False, error="Unsupported file format")
    
    content = await file.read()
    file_size_mb = len(content) / (1024 * 1024)
    if file_size_mb > MAX_FILE_SIZE_MB:
        return TranscriptionResponse(text="", success=False, error=f"File too large. Max {MAX_FILE_SIZE_MB}MB")
    
    transcription_id = str(uuid.uuid4())
    stop_event = threading.Event()
    stop_events[transcription_id] = stop_event
    
    result_container = {}
    
    thread = threading.Thread(target=run_transcription, args=(transcription_id, content, file.filename, result_container))
    thread.start()
    
    while thread.is_alive():
        await asyncio.sleep(0.5)
        if stop_event.is_set():
            break
    
    thread.join(timeout=1)
    
    if transcription_id in stop_events:
        del stop_events[transcription_id]
    
    if "result" not in result_container:
        return TranscriptionResponse(text="", success=False, error="Transcription timeout")
    
    result = result_container["result"]
    
    if result["success"]:
        transcription_progress.pop(transcription_id, None)
        return TranscriptionResponse(
            text=result["text"],
            success=True,
            audio_info=result["audio_info"],
            runtime_seconds=result["runtime_seconds"],
            runtime_formatted=result["runtime_formatted"]
        )
    else:
        transcription_progress.pop(transcription_id, None)
        return TranscriptionResponse(text="", success=False, error=result.get("error", "Unknown error"))


def run_batch_transcription(batch_id: str, files_data: list, model_name: str, save_docs: bool = False):
    model = get_model(model_name)
    file_results = []
    combined_text_parts = []
    total_start = time.time()

    for idx, (filename, content) in enumerate(files_data):
        file_start = time.time()
        temp_audio_path = None
        chunk_paths = []
        file_text = ""

        batch_progress[batch_id] = {
            "files_total": len(files_data),
            "files_done": idx,
            "current_file": filename,
            "current_chunk": 0,
            "total_chunks": 0,
            "status": "processing"
        }

        try:
            file_size_mb = len(content) / (1024 * 1024)
            suffix = Path(filename).suffix
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as f:
                f.write(content)
                temp_audio_path = f.name

            temp_audio_path = convert_to_wav(temp_audio_path)
            audio_info = get_audio_info(temp_audio_path, len(content))
            duration = audio_info["duration_seconds"]
            num_chunks = calculate_num_chunks(duration)

            batch_progress[batch_id]["total_chunks"] = num_chunks

            if num_chunks > 1:
                chunk_paths = split_audio_into_chunks(temp_audio_path, num_chunks)
                texts = []
                for ci, cp in enumerate(chunk_paths):
                    batch_progress[batch_id]["current_chunk"] = ci + 1
                    texts.append(model.transcribe(cp))
                file_text = " ".join(texts)
            else:
                batch_progress[batch_id]["current_chunk"] = 1
                file_text = model.transcribe(temp_audio_path)

            file_runtime = time.time() - file_start
            cpu_cores = torch.get_num_threads()
            save_log_entry(filename, duration, file_size_mb, file_runtime, num_chunks, cpu_cores, model_name)

            doc_path = None
            if save_docs and file_text:
                doc_basename = Path(filename).stem + ".doc"
                doc_path = str(OUTPUT_DIR / doc_basename)
                doc = Document()
                for line in file_text.split('\n'):
                    doc.add_paragraph(line)
                doc.save(doc_path)

            file_results.append(BatchFileResult(
                filename=filename, text=file_text, success=True,
                runtime_seconds=round(file_runtime, 2),
                doc_path=doc_path
            ))
        except Exception as e:
            logger.exception(f"Batch file {filename} failed")
            file_results.append(BatchFileResult(
                filename=filename, text="", success=False, error=str(e)
            ))
        finally:
            cleanup_files(chunk_paths + [temp_audio_path])

        if file_text:
            combined_text_parts.append(f"--- {filename} ---\n{file_text}")

    total_runtime = time.time() - total_start
    combined_text = "\n\n".join(combined_text_parts)

    batch_results[batch_id] = {
        "combined_text": combined_text,
        "files": file_results,
        "total_runtime": round(total_runtime, 2),
    }
    batch_progress[batch_id] = {
        "files_total": len(files_data),
        "files_done": len(files_data),
        "current_file": "",
        "current_chunk": 0,
        "total_chunks": 0,
        "status": "completed"
    }


@app.post("/api/transcribe-batch", response_model=BatchStartResponse)
async def transcribe_batch(files: List[UploadFile] = File(...), model: str = Form("GigaAM-v3"), save_docs: bool = Form(False)):
    if not files:
        return BatchStartResponse(batch_id="", success=False, error="No files provided")

    files_data = []
    for f in files:
        if not validate_file_format(f.filename):
            return BatchStartResponse(batch_id="", success=False, error=f"Unsupported format: {f.filename}")
        content = await f.read()
        file_size_mb = len(content) / (1024 * 1024)
        if file_size_mb > MAX_FILE_SIZE_MB:
            return BatchStartResponse(batch_id="", success=False, error=f"File too large: {f.filename}")
        files_data.append((f.filename, content))

    batch_id = str(uuid.uuid4())
    batch_progress[batch_id] = {
        "files_total": len(files_data),
        "files_done": 0,
        "current_file": "",
        "current_chunk": 0,
        "total_chunks": 0,
        "status": "starting"
    }

    thread = threading.Thread(target=run_batch_transcription, args=(batch_id, files_data, model, save_docs))
    thread.start()

    return BatchStartResponse(batch_id=batch_id, success=True)


@app.get("/api/batch-progress/{batch_id}")
async def get_batch_progress(batch_id: str):
    if batch_id in batch_progress:
        p = batch_progress[batch_id]
        return BatchProgressResponse(
            files_total=p["files_total"],
            files_done=p["files_done"],
            current_file=p["current_file"],
            current_chunk=p["current_chunk"],
            total_chunks=p["total_chunks"],
            status=p["status"]
        )
    return BatchProgressResponse(files_total=0, files_done=0, current_file="", current_chunk=0, total_chunks=0, status="not_found")


@app.get("/api/batch-result/{batch_id}")
async def get_batch_result(batch_id: str):
    if batch_id in batch_results:
        r = batch_results[batch_id]
        total_secs = r.get("total_runtime", 0)
        mins = int(total_secs // 60)
        secs = int(total_secs % 60)
        return BatchResultResponse(
            success=True,
            combined_text=r["combined_text"],
            files=r["files"],
            total_runtime_seconds=total_secs,
            total_runtime_formatted=f"{mins:02d}:{secs:02d}"
        )
    return BatchResultResponse(success=False, combined_text="", files=[], error="Batch not found")


@app.post("/api/save-doc", response_model=SaveResponse)
async def save_as_doc(request: SaveRequest):
    try:
        output_path = Path(request.output_path)
        if output_path.suffix.lower() != '.doc':
            output_path = output_path.with_suffix('.doc')
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        for paragraph_text in request.text.split('\n'):
            doc.add_paragraph(paragraph_text)
        doc.save(str(output_path))
        return SaveResponse(success=True, message="Document saved", file_path=str(output_path.resolve()))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/llm/enhance")
async def llm_enhance(req: LLMEnhanceRequest):
    try:
        model = req.model or "minimax-m3:cloud"
        async with httpx.AsyncClient(timeout=120) as client:
            resp = await client.post(
                f"{req.ollama_host}/api/chat",
                json={
                    "model": model,
                    "messages": [
                        {"role": "system", "content": req.prompt},
                        {"role": "user", "content": req.text}
                    ],
                    "stream": False
                }
            )
        if resp.status_code != 200:
            return {"success": False, "error": f"Ollama returned status {resp.status_code}: {resp.text}"}
        data = resp.json()
        return {"success": True, "result": data["message"]["content"]}
    except Exception as e:
        return {"success": False, "error": str(e)}


@app.get("/api/llm/models")
async def llm_models(host: str = "http://localhost:11434"):
    try:
        async with httpx.AsyncClient(timeout=5) as client:
            resp = await client.get(f"{host}/api/tags")
        if resp.status_code != 200:
            return {"success": False, "error": f"Ollama returned status {resp.status_code}", "models": []}
        data = resp.json()
        models = [m["name"] for m in data.get("models", [])]
        return {"success": True, "models": models}
    except Exception as e:
        return {"success": False, "error": str(e), "models": []}


@app.post("/api/llm/login")
async def llm_login():
    try:
        subprocess.Popen(
            ["ollama", "login"],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            creationflags=subprocess.CREATE_NO_WINDOW
        )
        return {"success": True, "message": "Browser opened for Ollama Cloud login. Complete auth in your browser, then pull cloud models."}
    except FileNotFoundError:
        return {"success": False, "message": "Ollama CLI not found. Install Ollama and add it to PATH."}
    except Exception as e:
        return {"success": False, "message": str(e)}


@app.post("/api/llm/pull")
async def llm_pull(model: str = Form("minimax-m3:cloud")):
    try:
        result = subprocess.run(
            ["ollama", "pull", model],
            capture_output=True, text=True, timeout=300
        )
        if result.returncode == 0:
            return {"success": True, "message": f"Model {model} pulled successfully"}
        else:
            return {"success": False, "message": result.stderr.strip() or f"Failed to pull {model}"}
    except FileNotFoundError:
        return {"success": False, "message": "Ollama CLI not found"}
    except subprocess.TimeoutExpired:
        return {"success": False, "message": f"Pulling {model} timed out (5 min)"}
    except Exception as e:
        return {"success": False, "message": str(e)}


@app.get("/api/health")
async def health_check():
    return {"status": "healthy"}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)